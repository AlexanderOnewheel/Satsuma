"""Every conversion, tool and action Satsuma can run, plus the registry the wheel reads.

An Op declares which file categories it applies to and a function
fn(paths, ctx) -> Result. Ops with ui=True are carried out by the app
itself (clipboard, folder pickers, floating shelf) and have no fn here.
"""
from __future__ import annotations

import base64
import io
import math
import os
import re
import shutil
import subprocess
import tarfile
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from . import media
from .media import AUDIO_PRESETS, GIF_FILTER, VIDEO_PRESETS, Runner
from .util import category, find_tool, human_size, sibling, unique_path

_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


# ---- framework --------------------------------------------------------------

@dataclass
class Param:
    key: str
    label: str
    default: str
    choices: list[str] | None = None
    hint: str = ""


@dataclass
class Result:
    outputs: list[Path] = field(default_factory=list)
    message: str = ""


class Ctx:
    """Passed to every op: parameters, progress, cancellation."""

    def __init__(self, params: dict | None = None,
                 progress: Callable[[float, str], None] | None = None) -> None:
        self.params = params or {}
        self._progress = progress
        self.runner = Runner()
        self.step = 0
        self.steps = 1

    def progress(self, frac: float, label: str = "") -> None:
        if self._progress:
            overall = (self.step + max(0.0, min(1.0, frac))) / max(1, self.steps)
            self._progress(overall, label)

    def check(self) -> None:
        if self.runner.cancelled:
            raise media.Cancelled()

    def cancel(self) -> None:
        self.runner.cancel()


@dataclass
class Op:
    id: str
    label: str
    glyph: str                     # short text shown in the wedge
    mode: str                      # "convert" | "tool" | "action"
    cats: frozenset[str]
    fn: Callable[[list[Path], Ctx], Result] | None = None
    params: list[Param] = field(default_factory=list)
    batch: bool = False            # True = one job over all files; False = per file
    min_files: int = 1
    ui: bool = False
    target_ext: str = ""           # for converts: hide if every file already has it

    def applies(self, paths: list[Path]) -> bool:
        if len(paths) < self.min_files:
            return False
        cats = {category(p) for p in paths}
        if "*" not in self.cats and not cats <= self.cats:
            return False
        if self.target_ext:
            exts = {_norm_ext(p) for p in paths}
            if exts == {self.target_ext}:
                return False
        return True


def _norm_ext(p: Path) -> str:
    e = p.suffix.lower().lstrip(".")
    return {"jpeg": "jpg", "tif": "tiff", "heif": "heic", "aif": "aiff"}.get(e, e)


REGISTRY: list[Op] = []


def op(**kw):
    def deco(fn):
        REGISTRY.append(Op(fn=fn, **kw))
        return fn
    return deco


def ops_for(paths: list[Path], mode: str) -> list[Op]:
    return [o for o in REGISTRY if o.mode == mode and o.applies(paths)]


def run_op(o: Op, paths: list[Path], ctx: Ctx) -> Result:
    """Run an op over the dropped files, per file or as one batch."""
    assert o.fn is not None
    if o.batch:
        ctx.steps = 1
        return o.fn(paths, ctx)
    total = Result()
    ctx.steps = len(paths)
    errors = []
    for i, p in enumerate(paths):
        ctx.step = i
        ctx.check()
        try:
            r = o.fn([p], ctx)
            total.outputs += r.outputs
            if r.message:
                total.message += (("\n\n" if total.message else "") +
                                  (f"{p.name}:\n" if len(paths) > 1 else "") + r.message)
        except media.Cancelled:
            raise
        except Exception as e:  # keep going on batch, report at end
            if len(paths) == 1:
                raise
            errors.append(f"{p.name}: {e}")
    if errors:
        total.message += ("\n\n" if total.message else "") + "Failed:\n" + "\n".join(errors)
    return total


# ---- shared parsing helpers ---------------------------------------------------

def parse_time(s: str) -> float:
    """'90', '1:30', '00:01:30.5' -> seconds."""
    s = s.strip()
    if not s:
        return 0.0
    parts = [float(x) for x in s.split(":")]
    secs = 0.0
    for x in parts:
        secs = secs * 60 + x
    return secs


def parse_target_size(s: str, original_bytes: int) -> int:
    """'8' / '8MB' / '500KB' / '50%' -> bytes."""
    s = s.strip().upper().replace(" ", "")
    if s.endswith("%"):
        return int(original_bytes * float(s[:-1]) / 100)
    mult = 1024 * 1024
    if s.endswith("KB"):
        mult, s = 1024, s[:-2]
    elif s.endswith("MB"):
        s = s[:-2]
    elif s.endswith("GB"):
        mult, s = 1024 ** 3, s[:-2]
    return int(float(s) * mult)


def atempo_chain(speed: float) -> str:
    """atempo only accepts 0.5–2.0 per stage, so chain stages for other speeds."""
    stages = []
    s = speed
    while s > 2.0:
        stages.append(2.0)
        s /= 2.0
    while s < 0.5:
        stages.append(0.5)
        s /= 0.5
    stages.append(s)
    return ",".join(f"atempo={x:.6g}" for x in stages)


# ---- images ----------------------------------------------------------------

IMAGE_FORMATS = [("jpg", "JPG"), ("png", "PNG"), ("heic", "HEIC"), ("webp", "WEBP"),
                 ("tiff", "TIFF"), ("bmp", "BMP"), ("gif", "GIF"), ("ico", "ICO"),
                 ("svg", "SVG"), ("pdf", "PDF")]

_heif_ready = False


def _pil():
    global _heif_ready
    from PIL import Image
    if not _heif_ready:
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            pass
        _heif_ready = True
    return Image


def open_image(path: Path):
    Image = _pil()
    if path.suffix.lower() == ".svg":
        import pymupdf
        doc = pymupdf.open(str(path))
        pix = doc[0].get_pixmap(dpi=192, alpha=True)
        return Image.open(io.BytesIO(pix.tobytes("png")))
    from PIL import ImageOps
    img = Image.open(path)
    img.load()
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    return img


def _flatten(img, bg=(255, 255, 255)):
    Image = _pil()
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        base = Image.new("RGB", img.size, bg)
        base.paste(img, mask=img.split()[-1])
        return base
    return img.convert("RGB") if img.mode != "RGB" else img


def save_image(img, out: Path, fmt: str, quality: int = 92, exif: bytes | None = None,
               keep_meta: bool = True) -> None:
    fmt = fmt.lower()
    kw: dict = {}
    if keep_meta and exif and fmt in ("jpg", "webp", "heic", "tiff"):
        kw["exif"] = exif
    if fmt == "jpg":
        _flatten(img).save(out, "JPEG", quality=quality, optimize=True, progressive=True, **kw)
    elif fmt == "png":
        (img if img.mode in ("RGB", "RGBA", "L", "LA", "P") else img.convert("RGBA")).save(
            out, "PNG", optimize=True)
    elif fmt == "webp":
        img.convert("RGBA" if "A" in img.mode else "RGB").save(out, "WEBP", quality=quality,
                                                               method=6, **kw)
    elif fmt == "heic":
        img.convert("RGBA" if "A" in img.mode else "RGB").save(out, "HEIF", quality=quality, **kw)
    elif fmt == "tiff":
        img.save(out, "TIFF", compression="tiff_lzw", **kw)
    elif fmt == "bmp":
        _flatten(img).save(out, "BMP")
    elif fmt == "gif":
        img.convert("RGBA" if "A" in img.mode else "RGB").save(out, "GIF")
    elif fmt == "ico":
        im = img.convert("RGBA")
        side = max(im.size)
        sq = _pil().new("RGBA", (side, side), (0, 0, 0, 0))
        sq.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        sizes = [s for s in (16, 24, 32, 48, 64, 128, 256) if s <= side] or [side]
        sq.save(out, "ICO", sizes=[(s, s) for s in sizes])
    elif fmt == "pdf":
        _flatten(img).save(out, "PDF", resolution=150.0)
    elif fmt == "svg":
        buf = io.BytesIO()
        img.convert("RGBA").save(buf, "PNG", optimize=True)
        b64 = base64.b64encode(buf.getvalue()).decode()
        w, h = img.size
        out.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" '
                       f'viewBox="0 0 {w} {h}"><image width="{w}" height="{h}" '
                       f'href="data:image/png;base64,{b64}"/></svg>', encoding="utf-8")
    else:
        raise ValueError(f"Unsupported image format {fmt}")


def _make_image_convert(ext: str, label: str):
    def fn(paths: list[Path], ctx: Ctx) -> Result:
        src = paths[0]
        img = open_image(src)
        out = sibling(src, ext)
        save_image(img, out, ext, exif=img.info.get("exif"))
        ctx.progress(1)
        return Result([out])
    REGISTRY.append(Op(id=f"img_to_{ext}", label=f"Convert to {label}", glyph=label,
                       mode="convert", cats=frozenset({"image"}), fn=fn, target_ext=ext))


for _ext, _label in IMAGE_FORMATS:
    _make_image_convert(_ext, _label)


@op(id="img_compress", label="Compress", glyph="Shrink", mode="tool", cats=frozenset({"image"}),
    params=[Param("level", "Strength", "Medium", ["Light", "Medium", "Strong"])])
def img_compress(paths, ctx):
    src = paths[0]
    q = {"Light": 82, "Medium": 70, "Strong": 50}[ctx.params.get("level", "Medium")]
    img = open_image(src)
    ext = _norm_ext(src)
    if ext not in ("jpg", "webp", "png", "heic"):
        ext = "jpg"
    out = sibling(src, ext, "compressed")
    if ext == "png":
        colors = {82: 256, 70: 192, 50: 96}[q]
        has_alpha = "A" in img.mode
        im = img.convert("RGBA") if has_alpha else img.convert("RGB")
        im.quantize(colors=colors, method=2 if has_alpha else 0).save(out, "PNG", optimize=True)
    else:
        save_image(img, out, ext, quality=q, exif=img.info.get("exif"))
    return Result([out], _saving(src, out))


@op(id="img_target", label="Compress to size", glyph="To Size", mode="tool",
    cats=frozenset({"image"}),
    params=[Param("size", "Target size", "500KB", hint="e.g. 500KB, 2MB or 40%")])
def img_target(paths, ctx):
    src = paths[0]
    target = parse_target_size(ctx.params.get("size", "500KB"), src.stat().st_size)
    img = open_image(src)
    ext = _norm_ext(src) if _norm_ext(src) in ("jpg", "webp") else "jpg"
    out = sibling(src, ext, "small")
    exif = img.info.get("exif")
    scale = 1.0
    for _ in range(6):  # shrink dimensions if even q=20 can't hit the target
        work = img if scale == 1.0 else img.resize(
            (max(1, int(img.width * scale)), max(1, int(img.height * scale))), _pil().LANCZOS)
        lo, hi, best = 20, 95, None
        while lo <= hi:
            mid = (lo + hi) // 2
            save_image(work, out, ext, quality=mid, exif=exif)
            size = out.stat().st_size
            if size <= target:
                best, lo = mid, mid + 1
            else:
                hi = mid - 1
            ctx.check()
        if best is not None:
            save_image(work, out, ext, quality=best, exif=exif)
            return Result([out], _saving(src, out))
        scale *= 0.75
    return Result([out], f"Couldn't reach {human_size(target)}; best effort "
                         f"{human_size(out.stat().st_size)}.")


@op(id="img_resize", label="Resize", glyph="Resize", mode="tool", cats=frozenset({"image"}),
    params=[Param("pct", "Scale (%) or max side (px)", "50%", hint="50%  or  1920px")])
def img_resize(paths, ctx):
    src = paths[0]
    img = open_image(src)
    v = ctx.params.get("pct", "50%").strip().lower()
    if v.endswith("px"):
        m = float(v[:-2])
        f = m / max(img.size)
    else:
        f = float(v.rstrip("%")) / 100
    new = img.resize((max(1, round(img.width * f)), max(1, round(img.height * f))),
                     _pil().LANCZOS)
    ext = _norm_ext(src) if _norm_ext(src) != "svg" else "png"
    out = sibling(src, ext, f"{new.width}x{new.height}")
    save_image(new, out, ext, exif=img.info.get("exif"))
    return Result([out])


@op(id="img_rotate", label="Rotate 90°", glyph="Rotate", mode="tool", cats=frozenset({"image"}))
def img_rotate(paths, ctx):
    src = paths[0]
    img = open_image(src).rotate(-90, expand=True)
    ext = _norm_ext(src) if _norm_ext(src) != "svg" else "png"
    out = sibling(src, ext, "rotated")
    save_image(img, out, ext, keep_meta=False)
    return Result([out])


@op(id="img_gray", label="Black & white", glyph="B&W", mode="tool", cats=frozenset({"image"}))
def img_gray(paths, ctx):
    src = paths[0]
    img = open_image(src)
    from PIL import ImageOps
    g = ImageOps.grayscale(_flatten(img) if "A" not in img.mode else img.convert("RGBA"))
    ext = _norm_ext(src) if _norm_ext(src) != "svg" else "png"
    out = sibling(src, ext, "bw")
    save_image(g, out, ext, exif=img.info.get("exif"))
    return Result([out])


@op(id="img_strip", label="Remove metadata", glyph="Strip", mode="tool",
    cats=frozenset({"image"}))
def img_strip(paths, ctx):
    src = paths[0]
    img = open_image(src)
    ext = _norm_ext(src) if _norm_ext(src) != "svg" else "png"
    clean = img.copy()
    clean.info = {}
    out = sibling(src, ext, "clean")
    save_image(clean, out, ext, quality=95, keep_meta=False)
    return Result([out], "Removed EXIF, GPS and camera data.")


@op(id="img_info", label="Show metadata", glyph="Info", mode="tool", cats=frozenset({"image"}))
def img_info(paths, ctx):
    src = paths[0]
    Image = _pil()
    from PIL.ExifTags import GPSTAGS, TAGS
    img = Image.open(src) if src.suffix.lower() != ".svg" else open_image(src)
    lines = [f"{img.width} × {img.height}  {img.mode}  {img.format or src.suffix[1:].upper()}",
             f"Size: {human_size(src.stat().st_size)}"]
    try:
        exif = img.getexif()
        for k, v in exif.items():
            name = TAGS.get(k, str(k))
            if isinstance(v, bytes):
                continue
            lines.append(f"{name}: {v}")
        gps = exif.get_ifd(0x8825)
        if gps:
            lines.append("GPS: " + ", ".join(f"{GPSTAGS.get(k, k)}={v}" for k, v in gps.items()))
    except Exception:
        pass
    if len(lines) == 2:
        lines.append("No EXIF metadata.")
    return Result([], "\n".join(lines))


@op(id="img_qr", label="Read QR code", glyph="QR", mode="tool", cats=frozenset({"image"}))
def img_qr(paths, ctx):
    try:
        import cv2
        import numpy as np
    except ImportError:
        raise RuntimeError("QR reading needs opencv-python-headless installed.")
    img = _flatten(open_image(paths[0]))
    arr = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
    det = cv2.QRCodeDetector()
    ok, texts, _, _ = det.detectAndDecodeMulti(arr)
    found = [t for t in (texts if ok else []) if t]
    if not found:
        t, _, _ = det.detectAndDecode(arr)
        found = [t] if t else []
    return Result([], "\n".join(found) if found else "No QR code found.")


@op(id="img_to_onepdf", label="Combine into one PDF", glyph="→PDF", mode="tool",
    cats=frozenset({"image"}), batch=True, min_files=2)
def img_to_onepdf(paths, ctx):
    imgs = [_flatten(open_image(p)) for p in sorted(paths, key=lambda p: p.name.lower())]
    out = unique_path(paths[0].with_name("Combined.pdf"))
    imgs[0].save(out, "PDF", save_all=True, append_images=imgs[1:], resolution=150.0)
    return Result([out])


@op(id="img_collage", label="Make collage", glyph="Collage", mode="tool",
    cats=frozenset({"image"}), batch=True, min_files=2)
def img_collage(paths, ctx):
    Image = _pil()
    imgs = [_flatten(open_image(p)) for p in sorted(paths, key=lambda p: p.name.lower())]
    cols = math.ceil(math.sqrt(len(imgs)))
    rows = math.ceil(len(imgs) / cols)
    cell = 800
    gap = 12
    sheet = Image.new("RGB", (cols * cell + (cols + 1) * gap, rows * cell + (rows + 1) * gap),
                      (255, 255, 255))
    for i, im in enumerate(imgs):
        im = im.copy()
        im.thumbnail((cell, cell), Image.LANCZOS)
        r, c = divmod(i, cols)
        x = gap + c * (cell + gap) + (cell - im.width) // 2
        y = gap + r * (cell + gap) + (cell - im.height) // 2
        sheet.paste(im, (x, y))
    out = unique_path(paths[0].with_name("Collage.jpg"))
    sheet.save(out, "JPEG", quality=90)
    return Result([out])


# ---- video -----------------------------------------------------------------

VIDEO_FORMATS = [("mp4", "MP4"), ("mov", "MOV"), ("mkv", "MKV"), ("webm", "WEBM"),
                 ("avi", "AVI"), ("wmv", "WMV"), ("gif", "GIF"), ("mp3", "MP3")]


def _make_video_convert(ext: str, label: str):
    def fn(paths, ctx):
        src = paths[0]
        dur = media.duration(src)
        out = sibling(src, ext)
        cb = lambda f: ctx.progress(f, label)  # noqa: E731
        if ext == "gif":
            ctx.runner.run(["-i", str(src), "-filter_complex",
                            GIF_FILTER.format(fps=12, w=640), "-loop", "0", str(out)], dur, cb)
        elif ext == "mp3":
            ctx.runner.run(["-i", str(src), "-vn", *AUDIO_PRESETS["mp3"], str(out)], dur, cb)
        elif ext == "mkv":
            try:  # mkv holds almost any codec, so a lossless remux usually works
                ctx.runner.run(["-i", str(src), "-map", "0", "-c", "copy", str(out)], dur, cb)
            except media.Cancelled:
                raise
            except Exception:
                ctx.runner.run(["-i", str(src), *VIDEO_PRESETS["mkv"], str(out)], dur, cb)
        else:
            ctx.runner.run(["-i", str(src), *VIDEO_PRESETS[ext], str(out)], dur, cb)
        return Result([out])
    REGISTRY.append(Op(id=f"vid_to_{ext}", label=f"Convert to {label}", glyph=label,
                       mode="convert", cats=frozenset({"video"}), fn=fn, target_ext=ext))


for _ext, _label in VIDEO_FORMATS:
    _make_video_convert(_ext, _label)


@op(id="vid_compress", label="Compress", glyph="Shrink", mode="tool", cats=frozenset({"video"}),
    params=[Param("level", "Strength", "Medium", ["Light", "Medium", "Strong"]),
            Param("max", "Max resolution", "Keep", ["Keep", "1080p", "720p", "480p"])])
def vid_compress(paths, ctx):
    src = paths[0]
    crf = {"Light": "24", "Medium": "28", "Strong": "32"}[ctx.params.get("level", "Medium")]
    vf = []
    mx = ctx.params.get("max", "Keep")
    if mx != "Keep":
        h = int(mx[:-1])
        vf = ["-vf", f"scale=-2:'min({h},ih)'"]
    out = sibling(src, "mp4", "compressed")
    ctx.runner.run(["-i", str(src), *vf, "-c:v", "libx264", "-preset", "slow", "-crf", crf,
                    "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k",
                    "-movflags", "+faststart", str(out)],
                   media.duration(src), lambda f: ctx.progress(f, "Compressing"))
    return Result([out], _saving(src, out))


def _two_pass_to_size(src: Path, out: Path, target: int, ctx: Ctx) -> None:
    dur = media.duration(src)
    if dur <= 0:
        raise RuntimeError("Couldn't read the video's duration.")
    has_audio = media.has_stream(src, "audio")
    audio_kbps = 96 if has_audio else 0
    total_kbps = target * 8 / 1024 / dur * 0.97  # 3% container headroom
    v_kbps = int(total_kbps - audio_kbps)
    if v_kbps < 50:
        raise RuntimeError(f"{human_size(target)} is too small for a {dur:.0f}s video.")
    with tempfile.TemporaryDirectory() as td:
        log = os.path.join(td, "pass")
        base = ["-i", str(src), "-c:v", "libx264", "-preset", "medium", "-b:v", f"{v_kbps}k",
                "-pix_fmt", "yuv420p", "-passlogfile", log]
        ctx.runner.run([*base, "-pass", "1", "-an", "-f", "mp4", os.devnull], dur,
                       lambda f: ctx.progress(f * 0.5, "Analyzing"))
        audio = ["-c:a", "aac", "-b:a", f"{audio_kbps}k"] if has_audio else ["-an"]
        ctx.runner.run([*base, "-pass", "2", *audio, "-movflags", "+faststart", str(out)], dur,
                       lambda f: ctx.progress(0.5 + f * 0.5, "Encoding"))


@op(id="vid_target", label="Compress to size", glyph="To Size", mode="tool",
    cats=frozenset({"video"}),
    params=[Param("size", "Target size", "25MB", hint="e.g. 25MB (Discord), 8MB, or 50%")])
def vid_target(paths, ctx):
    src = paths[0]
    target = parse_target_size(ctx.params.get("size", "25MB"), src.stat().st_size)
    out = sibling(src, "mp4", "small")
    _two_pass_to_size(src, out, target, ctx)
    return Result([out], _saving(src, out))


@op(id="vid_trim", label="Trim", glyph="Trim", mode="tool", cats=frozenset({"video", "audio"}),
    params=[Param("start", "Start", "0:00", hint="m:ss"),
            Param("end", "End (blank = to the end)", "", hint="m:ss")])
def vid_trim(paths, ctx):
    src = paths[0]
    start = parse_time(ctx.params.get("start", "0"))
    end_s = ctx.params.get("end", "").strip()
    dur = media.duration(src)
    end = parse_time(end_s) if end_s else dur
    if end <= start:
        raise ValueError("End must be after start.")
    is_video = category(src) == "video"
    ext = "mp4" if is_video else (_norm_ext(src) if _norm_ext(src) in AUDIO_PRESETS else "m4a")
    out = sibling(src, ext, "trimmed")
    codec = (VIDEO_PRESETS["mp4"] if is_video else AUDIO_PRESETS[ext])
    ctx.runner.run(["-ss", f"{start:.3f}", "-i", str(src), "-t", f"{end - start:.3f}",
                    *codec, str(out)], end - start, lambda f: ctx.progress(f, "Trimming"))
    return Result([out])


@op(id="vid_speed", label="Change speed", glyph="Speed", mode="tool",
    cats=frozenset({"video", "audio"}),
    params=[Param("speed", "Speed", "2x", ["0.25x", "0.5x", "0.75x", "1.25x", "1.5x", "2x",
                                             "4x", "8x"])])
def vid_speed(paths, ctx):
    src = paths[0]
    s = float(ctx.params.get("speed", "2x").rstrip("x"))
    dur = media.duration(src)
    is_video = category(src) == "video"
    has_audio = media.has_stream(src, "audio")
    tag = f"{s:g}x"
    if is_video:
        out = sibling(src, "mp4", tag)
        args = ["-i", str(src), "-filter:v", f"setpts=PTS/{s}"]
        args += ["-filter:a", atempo_chain(s)] if has_audio else ["-an"]
        args += VIDEO_PRESETS["mp4"]
    else:
        ext = _norm_ext(src) if _norm_ext(src) in AUDIO_PRESETS else "m4a"
        out = sibling(src, ext, tag)
        args = ["-i", str(src), "-filter:a", atempo_chain(s), *AUDIO_PRESETS[ext]]
    ctx.runner.run([*args, str(out)], dur / s, lambda f: ctx.progress(f, "Retiming"))
    return Result([out])


@op(id="vid_mute", label="Remove audio", glyph="Mute", mode="tool", cats=frozenset({"video"}))
def vid_mute(paths, ctx):
    src = paths[0]
    out = sibling(src, src.suffix, "muted")
    ctx.runner.run(["-i", str(src), "-map", "0:v", "-c", "copy", "-an", str(out)],
                   media.duration(src), ctx.progress)
    return Result([out])


@op(id="vid_snap", label="Snapshot", glyph="Snap", mode="tool", cats=frozenset({"video"}),
    params=[Param("at", "Time (blank = middle)", "", hint="m:ss")])
def vid_snap(paths, ctx):
    src = paths[0]
    at = ctx.params.get("at", "").strip()
    t = parse_time(at) if at else media.duration(src) / 2
    out = sibling(src, "png", f"frame-{int(t)}s")
    ctx.runner.run(["-ss", f"{t:.3f}", "-i", str(src), "-frames:v", "1", str(out)])
    return Result([out])


@op(id="vid_rotate", label="Rotate 90°", glyph="Rotate", mode="tool", cats=frozenset({"video"}))
def vid_rotate(paths, ctx):
    src = paths[0]
    out = sibling(src, "mp4", "rotated")
    ctx.runner.run(["-i", str(src), "-vf", "transpose=1", *VIDEO_PRESETS["mp4"], str(out)],
                   media.duration(src), ctx.progress)
    return Result([out])


@op(id="vid_crop", label="Crop to aspect", glyph="Crop", mode="tool", cats=frozenset({"video"}),
    params=[Param("aspect", "Aspect", "9:16", ["9:16", "1:1", "4:5", "16:9"])])
def vid_crop(paths, ctx):
    src = paths[0]
    a, b = (int(x) for x in ctx.params.get("aspect", "9:16").split(":"))
    vf = (f"crop='if(gt(iw/ih,{a}/{b}),trunc(ih*{a}/{b}/2)*2,iw)':"
          f"'if(gt(iw/ih,{a}/{b}),ih,trunc(iw*{b}/{a}/2)*2)'")
    out = sibling(src, "mp4", f"{a}x{b}")
    ctx.runner.run(["-i", str(src), "-vf", vf, *VIDEO_PRESETS["mp4"], str(out)],
                   media.duration(src), ctx.progress)
    return Result([out])


@op(id="vid_split", label="Split into parts", glyph="Split", mode="tool",
    cats=frozenset({"video", "audio"}),
    params=[Param("len", "Part length", "1:00", hint="m:ss per part")])
def vid_split(paths, ctx):
    src = paths[0]
    seg = parse_time(ctx.params.get("len", "1:00"))
    if seg <= 0:
        raise ValueError("Part length must be above zero.")
    folder = unique_path(src.with_name(src.stem + "-parts"))
    folder.mkdir()
    pattern = folder / f"{src.stem}-%03d{src.suffix}"
    ctx.runner.run(["-i", str(src), "-map", "0", "-c", "copy", "-f", "segment",
                    "-segment_time", f"{seg:.3f}", "-reset_timestamps", "1", str(pattern)],
                   media.duration(src), ctx.progress)
    return Result([folder])


@op(id="vid_join", label="Join", glyph="Join", mode="tool", cats=frozenset({"video"}),
    batch=True, min_files=2)
def vid_join(paths, ctx):
    files = sorted(paths, key=lambda p: p.name.lower())
    first = next(s for s in media.probe(files[0])["streams"] if s["codec_type"] == "video")
    W, H = int(first["width"]) // 2 * 2, int(first["height"]) // 2 * 2
    args, chains, labels = [], [], ""
    total = 0.0
    n = len(files)
    for i, f in enumerate(files):
        args += ["-i", str(f)]
    for i, f in enumerate(files):
        d = media.duration(f)
        total += d
        chains.append(f"[{i}:v]scale={W}:{H}:force_original_aspect_ratio=decrease,"
                      f"pad={W}:{H}:(ow-iw)/2:(oh-ih)/2,setsar=1,fps=30,format=yuv420p[v{i}]")
        if media.has_stream(f, "audio"):
            chains.append(f"[{i}:a]aresample=48000,aformat=channel_layouts=stereo[a{i}]")
        else:
            chains.append(f"anullsrc=r=48000:cl=stereo,atrim=0:{d:.3f}[a{i}]")
        labels += f"[v{i}][a{i}]"
    chains.append(f"{labels}concat=n={n}:v=1:a=1[v][a]")
    out = unique_path(files[0].with_name("Joined.mp4"))
    ctx.runner.run([*args, "-filter_complex", ";".join(chains), "-map", "[v]", "-map", "[a]",
                    "-c:v", "libx264", "-crf", "20", "-preset", "medium", "-c:a", "aac",
                    "-b:a", "160k", "-movflags", "+faststart", str(out)], total, ctx.progress)
    return Result([out])


@op(id="av_strip", label="Remove metadata", glyph="Strip", mode="tool",
    cats=frozenset({"video", "audio"}))
def av_strip(paths, ctx):
    src = paths[0]
    out = sibling(src, src.suffix, "clean")
    ctx.runner.run(["-i", str(src), "-map", "0", "-map_metadata", "-1", "-map_chapters", "-1",
                    "-c", "copy", str(out)], media.duration(src), ctx.progress)
    return Result([out], "Removed tags, GPS, device and chapter metadata.")


@op(id="av_info", label="Show metadata", glyph="Info", mode="tool",
    cats=frozenset({"video", "audio"}))
def av_info(paths, ctx):
    info = media.probe(paths[0])
    fmt = info.get("format", {})
    lines = [f"Container: {fmt.get('format_long_name', '?')}",
             f"Duration: {float(fmt.get('duration', 0)):.2f}s   "
             f"Size: {human_size(int(fmt.get('size', 0)))}   "
             f"Bitrate: {int(fmt.get('bit_rate', 0)) // 1000} kb/s"]
    for s in info.get("streams", []):
        t = s.get("codec_type")
        if t == "video":
            lines.append(f"Video: {s.get('codec_name')} {s.get('width')}×{s.get('height')} "
                         f"@ {s.get('r_frame_rate')} fps")
        elif t == "audio":
            lines.append(f"Audio: {s.get('codec_name')} {s.get('sample_rate')} Hz "
                         f"{s.get('channels')} ch")
    for k, v in (fmt.get("tags") or {}).items():
        lines.append(f"{k}: {v}")
    return Result([], "\n".join(lines))


# ---- audio -----------------------------------------------------------------

AUDIO_FORMATS = [("mp3", "MP3"), ("m4a", "M4A"), ("wav", "WAV"), ("flac", "FLAC"),
                 ("ogg", "OGG"), ("opus", "OPUS"), ("aiff", "AIFF")]


def _make_audio_convert(ext: str, label: str):
    def fn(paths, ctx):
        src = paths[0]
        out = sibling(src, ext)
        ctx.runner.run(["-i", str(src), "-vn", *AUDIO_PRESETS[ext], str(out)],
                       media.duration(src), lambda f: ctx.progress(f, label))
        return Result([out])
    REGISTRY.append(Op(id=f"aud_to_{ext}", label=f"Convert to {label}", glyph=label,
                       mode="convert", cats=frozenset({"audio"}), fn=fn, target_ext=ext))


for _ext, _label in AUDIO_FORMATS:
    _make_audio_convert(_ext, _label)


def _audio_out(src: Path, tag: str) -> tuple[Path, list[str]]:
    ext = _norm_ext(src) if _norm_ext(src) in AUDIO_PRESETS else "m4a"
    return sibling(src, ext, tag), AUDIO_PRESETS[ext]


@op(id="aud_norm", label="Normalize loudness", glyph="Level", mode="tool",
    cats=frozenset({"audio"}),
    params=[Param("lufs", "Target", "-16 LUFS (podcast)",
                  ["-14 LUFS (streaming)", "-16 LUFS (podcast)", "-23 LUFS (broadcast)"])])
def aud_norm(paths, ctx):
    src = paths[0]
    lufs = ctx.params.get("lufs", "-16").split()[0]
    out, codec = _audio_out(src, "normalized")
    ctx.runner.run(["-i", str(src), "-af", f"loudnorm=I={lufs}:TP=-1.5:LRA=11", "-ar", "48000",
                    *codec, str(out)], media.duration(src), ctx.progress)
    return Result([out])


@op(id="aud_channels", label="Mono / stereo", glyph="Chan", mode="tool",
    cats=frozenset({"audio"}),
    params=[Param("ch", "Output", "Mono", ["Mono", "Stereo", "Left only", "Right only"])])
def aud_channels(paths, ctx):
    src = paths[0]
    ch = ctx.params.get("ch", "Mono")
    af = {"Mono": ["-ac", "1"], "Stereo": ["-ac", "2"],
          "Left only": ["-af", "pan=mono|c0=c0"], "Right only": ["-af", "pan=mono|c0=c1"]}[ch]
    out, codec = _audio_out(src, ch.lower().replace(" ", "-"))
    ctx.runner.run(["-i", str(src), *af, *codec, str(out)], media.duration(src), ctx.progress)
    return Result([out])


@op(id="aud_bleep", label="Bleep sections", glyph="Bleep", mode="tool",
    cats=frozenset({"audio", "video"}),
    params=[Param("ranges", "Ranges to bleep", "0:05-0:06",
                  hint="comma separated, e.g. 0:05-0:06, 1:12.5-1:13.2")])
def aud_bleep(paths, ctx):
    src = paths[0]
    ranges = []
    for part in ctx.params.get("ranges", "").split(","):
        if "-" in part:
            a, b = part.split("-", 1)
            ranges.append((parse_time(a), parse_time(b)))
    if not ranges:
        raise ValueError("Enter at least one range like 0:05-0:06.")
    dur = media.duration(src)
    cond = "+".join(f"between(t,{a:.3f},{b:.3f})" for a, b in ranges)
    fc = (f"[0:a]volume=enable='{cond}':volume=0[muted];"
          f"sine=f=1000:sample_rate=48000:d={dur:.3f},volume=0.35,"
          f"volume=enable='not({cond})':volume=0[beep];"
          f"[muted][beep]amix=inputs=2:duration=first:normalize=0[a]")
    if category(src) == "video":
        out = sibling(src, "mp4", "bleeped")
        args = ["-i", str(src), "-filter_complex", fc, "-map", "0:v", "-map", "[a]",
                "-c:v", "copy", "-c:a", "aac", "-b:a", "192k"]
    else:
        out, codec = _audio_out(src, "bleeped")
        args = ["-i", str(src), "-filter_complex", fc, "-map", "[a]", *codec]
    ctx.runner.run([*args, str(out)], dur, ctx.progress)
    return Result([out])


@op(id="aud_target", label="Compress to size", glyph="To Size", mode="tool",
    cats=frozenset({"audio"}), params=[Param("size", "Target size", "5MB", hint="5MB or 50%")])
def aud_target(paths, ctx):
    src = paths[0]
    target = parse_target_size(ctx.params.get("size", "5MB"), src.stat().st_size)
    dur = media.duration(src)
    kbps = max(24, min(320, int(target * 8 / 1024 / max(dur, 0.1) * 0.97)))
    out = sibling(src, "m4a", "small")
    ctx.runner.run(["-i", str(src), "-vn", "-c:a", "aac", "-b:a", f"{kbps}k", str(out)], dur,
                   ctx.progress)
    return Result([out], _saving(src, out))


@op(id="aud_join", label="Join", glyph="Join", mode="tool", cats=frozenset({"audio"}),
    batch=True, min_files=2)
def aud_join(paths, ctx):
    files = sorted(paths, key=lambda p: p.name.lower())
    args, labels = [], ""
    for i, f in enumerate(files):
        args += ["-i", str(f)]
        labels += f"[a{i}]"
    pre = ";".join(f"[{i}:a]aresample=48000,aformat=channel_layouts=stereo[a{i}]"
                   for i in range(len(files)))
    fc = f"{pre};{labels}concat=n={len(files)}:v=0:a=1[a]"
    ext = _norm_ext(files[0]) if _norm_ext(files[0]) in AUDIO_PRESETS else "m4a"
    out = unique_path(files[0].with_name(f"Joined.{ext}"))
    total = sum(media.duration(f) for f in files)
    ctx.runner.run([*args, "-filter_complex", fc, "-map", "[a]", *AUDIO_PRESETS[ext], str(out)],
                   total, ctx.progress)
    return Result([out])


# ---- PDF -------------------------------------------------------------------

def _pymupdf():
    import pymupdf
    return pymupdf


def _render_pages(src: Path, ext: str, ctx: Ctx, dpi: int = 150) -> list[Path]:
    doc = _pymupdf().open(str(src))
    if doc.page_count == 1:
        targets = [sibling(src, ext)]
    else:
        folder = unique_path(src.with_name(src.stem + "-pages"))
        folder.mkdir()
        targets = [folder / f"{src.stem}-{i + 1:03d}.{ext}" for i in range(doc.page_count)]
    for i, page in enumerate(doc):
        ctx.check()
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        if ext == "png":
            pix.save(str(targets[i]))
        else:
            from PIL import Image
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            save_image(im, targets[i], ext, quality=90)
        ctx.progress((i + 1) / doc.page_count, f"Page {i + 1}")
    return [targets[0]] if doc.page_count == 1 else [targets[0].parent]


for _ext, _label in (("png", "PNG"), ("jpg", "JPG"), ("webp", "WEBP")):
    def _mk(ext=_ext, label=_label):
        REGISTRY.append(Op(id=f"pdf_to_{ext}", label=f"Pages to {label}", glyph=label,
                           mode="convert", cats=frozenset({"pdf"}),
                           fn=lambda paths, ctx: Result(_render_pages(paths[0], ext, ctx))))
    _mk()


@op(id="pdf_to_txt", label="Extract text", glyph="TXT", mode="convert", cats=frozenset({"pdf"}))
def pdf_to_txt(paths, ctx):
    src = paths[0]
    doc = _pymupdf().open(str(src))
    text = "\n\f".join(p.get_text() for p in doc)
    out = sibling(src, "txt")
    out.write_text(text, encoding="utf-8")
    return Result([out], "" if text.strip() else
                  "No text layer found; this PDF may be a scan.")


@op(id="pdf_to_docx", label="To Word (text)", glyph="DOCX", mode="convert",
    cats=frozenset({"pdf"}))
def pdf_to_docx(paths, ctx):
    import docx
    src = paths[0]
    doc = _pymupdf().open(str(src))
    d = docx.Document()
    for i, page in enumerate(doc):
        if i:
            d.add_page_break()
        for block in page.get_text("blocks"):
            if block[6] == 0 and block[4].strip():
                d.add_paragraph(" ".join(block[4].split()))
    out = sibling(src, "docx")
    d.save(out)
    return Result([out], "Text and paragraphs carried over; layout and images are not.")


@op(id="pdf_compress", label="Compress", glyph="Shrink", mode="tool", cats=frozenset({"pdf"}),
    params=[Param("level", "Strength", "Medium",
                  ["Light (200 dpi)", "Medium (150 dpi)", "Strong (96 dpi)"])])
def pdf_compress(paths, ctx):
    src = paths[0]
    level = ctx.params.get("level", "Medium")
    dpi, q = (200, 80) if level.startswith("Light") else (96, 55) if level.startswith(
        "Strong") else (150, 70)
    doc = _pymupdf().open(str(src))
    try:
        doc.rewrite_images(dpi_threshold=dpi + 20, dpi_target=dpi, quality=q)
    except Exception:
        pass  # older PyMuPDF: fall back to lossless clean-up only
    ctx.progress(0.7, "Saving")
    out = sibling(src, "pdf", "compressed")
    doc.save(str(out), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
             clean=True, use_objstms=1)
    if out.stat().st_size >= src.stat().st_size:
        out.unlink()
        return Result([], "Already well compressed; no smaller version was possible.")
    return Result([out], _saving(src, out))


@op(id="pdf_merge", label="Merge", glyph="Merge", mode="tool", cats=frozenset({"pdf", "image"}),
    batch=True, min_files=2)
def pdf_merge(paths, ctx):
    fitz = _pymupdf()
    files = sorted(paths, key=lambda p: p.name.lower())
    if all(category(f) == "image" for f in files):
        return img_to_onepdf(files, ctx)
    out_doc = fitz.open()
    for i, f in enumerate(files):
        if category(f) == "pdf":
            out_doc.insert_pdf(fitz.open(str(f)))
        else:
            img = fitz.open(str(f))
            out_doc.insert_pdf(fitz.open("pdf", img.convert_to_pdf()))
        ctx.progress((i + 1) / len(files))
    out = unique_path(files[0].with_name("Merged.pdf"))
    out_doc.save(str(out), garbage=3, deflate=True)
    return Result([out])


@op(id="pdf_split", label="Split pages", glyph="Split", mode="tool", cats=frozenset({"pdf"}),
    params=[Param("ranges", "Pages", "each",
                  hint="'each' for one file per page, or ranges like 1-3, 4, 5-9")])
def pdf_split(paths, ctx):
    fitz = _pymupdf()
    src = paths[0]
    doc = fitz.open(str(src))
    spec = ctx.params.get("ranges", "each").strip().lower()
    if spec in ("", "each"):
        groups = [(i, i) for i in range(doc.page_count)]
    else:
        groups = []
        for part in spec.split(","):
            a, _, b = part.strip().partition("-")
            lo, hi = int(a) - 1, int(b or a) - 1
            if not 0 <= lo <= hi < doc.page_count:
                raise ValueError(f"Page range {part.strip()} is outside 1-{doc.page_count}.")
            groups.append((lo, hi))
    folder = unique_path(src.with_name(src.stem + "-split"))
    folder.mkdir()
    for n, (lo, hi) in enumerate(groups):
        part = fitz.open()
        part.insert_pdf(doc, from_page=lo, to_page=hi)
        name = f"{src.stem}-p{lo + 1}" + (f"-{hi + 1}" if hi > lo else "") + ".pdf"
        part.save(str(folder / name))
        ctx.progress((n + 1) / len(groups))
    return Result([folder])


@op(id="pdf_rotate", label="Rotate 90°", glyph="Rotate", mode="tool", cats=frozenset({"pdf"}))
def pdf_rotate(paths, ctx):
    src = paths[0]
    doc = _pymupdf().open(str(src))
    for p in doc:
        p.set_rotation((p.rotation + 90) % 360)
    out = sibling(src, "pdf", "rotated")
    doc.save(str(out))
    return Result([out])


@op(id="pdf_strip", label="Remove metadata", glyph="Strip", mode="tool", cats=frozenset({"pdf"}))
def pdf_strip(paths, ctx):
    src = paths[0]
    doc = _pymupdf().open(str(src))
    doc.set_metadata({})
    doc.del_xml_metadata()
    out = sibling(src, "pdf", "clean")
    doc.save(str(out), garbage=3, deflate=True)
    return Result([out], "Removed author, title, producer and XMP metadata.")


@op(id="pdf_info", label="Show metadata", glyph="Info", mode="tool", cats=frozenset({"pdf"}))
def pdf_info(paths, ctx):
    src = paths[0]
    doc = _pymupdf().open(str(src))
    lines = [f"Pages: {doc.page_count}   Size: {human_size(src.stat().st_size)}",
             f"Encrypted: {'yes' if doc.is_encrypted else 'no'}"]
    lines += [f"{k}: {v}" for k, v in (doc.metadata or {}).items() if v]
    return Result([], "\n".join(lines))


@op(id="pdf_images", label="Extract images", glyph="Images", mode="tool",
    cats=frozenset({"pdf"}))
def pdf_images(paths, ctx):
    src = paths[0]
    doc = _pymupdf().open(str(src))
    folder = unique_path(src.with_name(src.stem + "-images"))
    folder.mkdir()
    n = 0
    seen = set()
    for pno, page in enumerate(doc):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            data = doc.extract_image(xref)
            n += 1
            (folder / f"{src.stem}-{pno + 1:03d}-{n:03d}.{data['ext']}").write_bytes(data["image"])
    if n == 0:
        folder.rmdir()
        return Result([], "No embedded images found.")
    return Result([folder], f"Extracted {n} images.")


# ---- documents ---------------------------------------------------------------

_SOFFICE_DIRS = [r"%ProgramFiles%\LibreOffice\program", r"%ProgramFiles(x86)%\LibreOffice\program"]
_OFFICE_EXT = {".docx", ".doc", ".rtf", ".odt", ".pptx", ".ppt", ".odp", ".xlsx", ".xls", ".ods"}
_MUPDF_DOC_EXT = {".epub", ".xps", ".mobi", ".fb2", ".cbz"}


def _office_to(src: Path, out: Path, fmt: str = "pdf") -> None:
    soffice = find_tool("soffice", _SOFFICE_DIRS)
    if soffice:
        with tempfile.TemporaryDirectory() as td:
            r = subprocess.run([soffice, "--headless", "--convert-to", fmt, "--outdir", td,
                                str(src)], capture_output=True, text=True, timeout=300,
                               creationflags=_NO_WINDOW)
            produced = list(Path(td).glob(f"*.{fmt}"))
            if r.returncode != 0 or not produced:
                raise RuntimeError(r.stderr.strip() or "LibreOffice conversion failed")
            shutil.move(str(produced[0]), out)
        return
    if os.name == "nt" and fmt == "pdf":
        ext = src.suffix.lower()
        s, o = str(src).replace("'", "''"), str(out).replace("'", "''")
        if ext in (".pptx", ".ppt", ".odp"):
            ps = (f"$a=New-Object -ComObject PowerPoint.Application;"
                  f"$d=$a.Presentations.Open('{s}',$true,$false,$false);"
                  f"$d.SaveAs('{o}',32);$d.Close();$a.Quit()")
        elif ext in (".xlsx", ".xls", ".ods"):
            ps = (f"$a=New-Object -ComObject Excel.Application;$a.DisplayAlerts=$false;"
                  f"$d=$a.Workbooks.Open('{s}',0,$true);$d.ExportAsFixedFormat(0,'{o}');"
                  f"$d.Close($false);$a.Quit()")
        else:
            ps = (f"$a=New-Object -ComObject Word.Application;$a.Visible=$false;"
                  f"$d=$a.Documents.Open('{s}',$false,$true);$d.SaveAs([ref]'{o}',[ref]17);"
                  f"$d.Close([ref]0);$a.Quit()")
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                           capture_output=True, text=True, timeout=300, creationflags=_NO_WINDOW)
        if r.returncode == 0 and out.exists():
            return
        raise RuntimeError("Office conversion failed. Install LibreOffice or Microsoft Office. "
                           + r.stderr.strip()[:200])
    raise RuntimeError("Converting Office files needs LibreOffice (or Microsoft Office).")


def _text_of(src: Path) -> str:
    ext = src.suffix.lower()
    if ext == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(str(src)).paragraphs)
    if ext in _MUPDF_DOC_EXT:
        return "\n".join(p.get_text() for p in _pymupdf().open(str(src)))
    return src.read_text(encoding="utf-8", errors="replace")


def _text_to_pdf(text: str, out: Path, markdown: bool = False) -> None:
    fitz = _pymupdf()
    import html as _html
    if markdown:
        body = _md_to_html(text)
    else:
        body = "<pre style='font-family:monospace;white-space:pre-wrap'>" + \
            _html.escape(text) + "</pre>"
    story = fitz.Story(html=f"<body style='font-family:sans-serif;font-size:11pt'>{body}</body>")
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def _md_to_html(md: str) -> str:
    import html as _html
    out, in_list = [], False
    for line in md.splitlines():
        esc = _html.escape(line)
        esc = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)
        esc = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"<i>\1</i>", esc)
        esc = re.sub(r"`(.+?)`", r"<code>\1</code>", esc)
        m = re.match(r"^(#{1,6})\s+(.*)", esc)
        if re.match(r"^\s*[-*]\s+", line):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append("<li>" + re.sub(r"^\s*[-*]\s+", "", esc) + "</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if m:
            n = len(m.group(1))
            out.append(f"<h{n}>{m.group(2)}</h{n}>")
        elif esc.strip():
            out.append(f"<p>{esc}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


@op(id="doc_to_pdf", label="Convert to PDF", glyph="PDF", mode="convert",
    cats=frozenset({"document"}))
def doc_to_pdf(paths, ctx):
    src = paths[0]
    ext = src.suffix.lower()
    out = sibling(src, "pdf")
    if ext in _MUPDF_DOC_EXT:
        fitz = _pymupdf()
        fitz.open("pdf", fitz.open(str(src)).convert_to_pdf()).save(str(out), garbage=3,
                                                                    deflate=True)
    elif ext in (".txt", ".md"):
        _text_to_pdf(src.read_text(encoding="utf-8", errors="replace"), out, ext == ".md")
    elif ext in _OFFICE_EXT:
        ctx.progress(0.2, "Opening in Office")
        _office_to(src, out, "pdf")
    else:
        raise RuntimeError(f"Can't convert {ext} to PDF.")
    return Result([out])


@op(id="doc_to_txt", label="Convert to TXT", glyph="TXT", mode="convert",
    cats=frozenset({"document"}), target_ext="txt")
def doc_to_txt(paths, ctx):
    src = paths[0]
    if src.suffix.lower() in _OFFICE_EXT - {".docx"}:
        out = sibling(src, "txt")
        _office_to(src, out, "txt")
        return Result([out])
    out = sibling(src, "txt")
    out.write_text(_text_of(src), encoding="utf-8")
    return Result([out])


@op(id="doc_to_docx", label="Convert to DOCX", glyph="DOCX", mode="convert",
    cats=frozenset({"document"}), target_ext="docx")
def doc_to_docx(paths, ctx):
    import docx
    src = paths[0]
    out = sibling(src, "docx")
    if src.suffix.lower() in (".doc", ".rtf", ".odt"):
        _office_to(src, out, "docx")
        return Result([out])
    d = docx.Document()
    for line in _text_of(src).splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line) if src.suffix.lower() == ".md" else None
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        else:
            d.add_paragraph(line)
    d.save(out)
    return Result([out])


@op(id="doc_to_epub", label="Convert to EPUB", glyph="EPUB", mode="convert",
    cats=frozenset({"document", "pdf"}), target_ext="epub")
def doc_to_epub(paths, ctx):
    """Minimal reflowable EPUB 3 from the document's text."""
    import html as _html
    import uuid
    src = paths[0]
    if src.suffix.lower() == ".pdf":
        doc = _pymupdf().open(str(src))
        chapters = [[" ".join(b[4].split()) for b in p.get_text("blocks")
                     if b[6] == 0 and b[4].strip()] for p in doc]
    else:
        text = _text_of(src)
        chapters = [[l for l in text.split("\n\n") if l.strip()]]
    title = src.stem
    out = sibling(src, "epub")
    uid = str(uuid.uuid4())
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:'
                   'opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/c.opf" '
                   'media-type="application/oebps-package+xml"/></rootfiles></container>')
        items, spine, nav = [], [], []
        for i, paras in enumerate(chapters):
            body = "".join(f"<p>{_html.escape(p)}</p>" for p in paras)
            z.writestr(f"OEBPS/ch{i}.xhtml",
                       f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/'
                       f'1999/xhtml"><head><title>{_html.escape(title)}</title></head>'
                       f'<body>{body}</body></html>')
            items.append(f'<item id="c{i}" href="ch{i}.xhtml" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{i}"/>')
            nav.append(f'<li><a href="ch{i}.xhtml">Part {i + 1}</a></li>')
        z.writestr("OEBPS/nav.xhtml",
                   '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/'
                   'xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents'
                   '</title></head><body><nav epub:type="toc"><ol>' + "".join(nav) +
                   '</ol></nav></body></html>')
        z.writestr("OEBPS/c.opf",
                   f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/'
                   f'2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http:'
                   f'//purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:{uid}'
                   f'</dc:identifier><dc:title>{_html.escape(title)}</dc:title><dc:language>en'
                   f'</dc:language><meta property="dcterms:modified">2026-01-01T00:00:00Z</meta>'
                   f'</metadata><manifest><item id="nav" href="nav.xhtml" media-type="'
                   f'application/xhtml+xml" properties="nav"/>{"".join(items)}</manifest>'
                   f'<spine>{"".join(spine)}</spine></package>')
    return Result([out], "Text-only EPUB; images and layout are not carried over.")


# ---- subtitles ---------------------------------------------------------------

@op(id="sub_to_vtt", label="Convert to VTT", glyph="VTT", mode="convert",
    cats=frozenset({"subtitle"}), target_ext="vtt")
def sub_to_vtt(paths, ctx):
    src = paths[0]
    text = src.read_text(encoding="utf-8-sig", errors="replace")
    body = re.sub(r"(\d\d:\d\d:\d\d),(\d\d\d)", r"\1.\2", text)
    out = sibling(src, "vtt")
    out.write_text("WEBVTT\n\n" + body.strip() + "\n", encoding="utf-8")
    return Result([out])


@op(id="sub_to_srt", label="Convert to SRT", glyph="SRT", mode="convert",
    cats=frozenset({"subtitle"}), target_ext="srt")
def sub_to_srt(paths, ctx):
    src = paths[0]
    text = src.read_text(encoding="utf-8-sig", errors="replace")
    blocks = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    cues = []
    for b in blocks:
        lines = [l for l in b.strip().split("\n") if l.strip()]
        idx = next((i for i, l in enumerate(lines) if "-->" in l), None)
        if idx is None:
            continue
        timing = lines[idx]
        def fix(t: str) -> str:
            t = t.strip().split(" ")[0].replace(".", ",")
            return t if t.count(":") == 2 else "00:" + t
        a, b2 = timing.split("-->")
        cues.append(f"{len(cues) + 1}\n{fix(a)} --> {fix(b2)}\n" + "\n".join(lines[idx + 1:]))
    out = sibling(src, "srt")
    out.write_text("\n\n".join(cues) + "\n", encoding="utf-8")
    return Result([out])


# ---- archives ----------------------------------------------------------------

_7Z_DIRS = [r"%ProgramFiles%\7-Zip", r"%ProgramFiles(x86)%\7-Zip"]


def _extract(src: Path, dest: Path) -> None:
    name = src.name.lower()
    if name.endswith(".zip"):
        with zipfile.ZipFile(src) as z:
            for m in z.infolist():
                target = (dest / m.filename).resolve()
                if not str(target).startswith(str(dest.resolve())):
                    raise RuntimeError("Archive contains unsafe paths.")
            z.extractall(dest)
    elif name.endswith((".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tar.xz")):
        with tarfile.open(src) as t:
            t.extractall(dest, filter="data")
    elif name.endswith((".gz", ".bz2", ".xz")) and not name.endswith(".tar.gz"):
        import bz2
        import gzip
        import lzma
        opener = {".gz": gzip.open, ".bz2": bz2.open, ".xz": lzma.open}[src.suffix.lower()]
        with opener(src, "rb") as f_in, open(dest / src.stem, "wb") as f_out:
            shutil.copyfileobj(f_in, f_out)
    else:  # 7z, rar and anything else 7-Zip understands
        sz = find_tool("7z", _7Z_DIRS) or find_tool("7za", _7Z_DIRS)
        if sz:
            r = subprocess.run([sz, "x", "-y", f"-o{dest}", str(src)], capture_output=True,
                               text=True, creationflags=_NO_WINDOW)
            if r.returncode != 0:
                raise RuntimeError(r.stderr.strip() or "7-Zip failed")
        elif os.name == "nt":  # Windows 11 tar.exe (libarchive) reads rar/7z
            r = subprocess.run(["tar", "-xf", str(src), "-C", str(dest)], capture_output=True,
                               text=True, creationflags=_NO_WINDOW)
            if r.returncode != 0:
                raise RuntimeError("Install 7-Zip to open this archive.")
        else:
            raise RuntimeError("Install 7-Zip to open this archive.")


def _archive_stem(src: Path) -> str:
    n = src.name
    for suf in (".tar.gz", ".tar.bz2", ".tar.xz", ".tgz"):
        if n.lower().endswith(suf):
            return n[: -len(suf)]
    return src.stem


@op(id="arc_extract", label="Extract here", glyph="Unzip", mode="convert",
    cats=frozenset({"archive"}))
def arc_extract(paths, ctx):
    src = paths[0]
    dest = unique_path(src.with_name(_archive_stem(src)))
    dest.mkdir()
    _extract(src, dest)
    kids = list(dest.iterdir())
    if len(kids) == 1 and kids[0].is_dir() and not (src.parent / kids[0].name).exists():
        final = src.parent / kids[0].name  # avoid folder-in-folder
        shutil.move(str(kids[0]), final)
        dest.rmdir()
        dest = final
    return Result([dest])


def _pack(sources: list[Path], out: Path, fmt: str) -> None:
    if fmt == "zip":
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as z:
            for s in sources:
                if s.is_dir():
                    for f in s.rglob("*"):
                        z.write(f, f.relative_to(s.parent))
                else:
                    z.write(s, s.name)
    else:
        with tarfile.open(out, "w:gz") as t:
            for s in sources:
                t.add(s, arcname=s.name)


@op(id="arc_to_zip", label="Convert to ZIP", glyph="ZIP", mode="convert",
    cats=frozenset({"archive"}), target_ext="zip")
def arc_to_zip(paths, ctx):
    src = paths[0]
    with tempfile.TemporaryDirectory() as td:
        _extract(src, Path(td))
        out = unique_path(src.with_name(_archive_stem(src) + ".zip"))
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for f in Path(td).rglob("*"):
                z.write(f, f.relative_to(td))
    return Result([out])


@op(id="arc_to_tgz", label="Convert to TAR.GZ", glyph="TAR.GZ", mode="convert",
    cats=frozenset({"archive"}))
def arc_to_tgz(paths, ctx):
    src = paths[0]
    if src.name.lower().endswith((".tar.gz", ".tgz")):
        raise ValueError("Already a .tar.gz")
    with tempfile.TemporaryDirectory() as td:
        _extract(src, Path(td))
        out = unique_path(src.with_name(_archive_stem(src) + ".tar.gz"))
        with tarfile.open(out, "w:gz") as t:
            for f in Path(td).iterdir():
                t.add(f, arcname=f.name)
    return Result([out])


# ---- actions (work on anything) ------------------------------------------------

@op(id="act_zip", label="Compress to ZIP", glyph="ZIP", mode="action", cats=frozenset({"*"}),
    batch=True)
def act_zip(paths, ctx):
    name = (paths[0].name if paths[0].is_dir() else paths[0].stem) if len(paths) == 1 \
        else "Archive"
    out = unique_path(paths[0].with_name(name + ".zip"))
    _pack(paths, out, "zip")
    return Result([out])


for _id, _label, _glyph in (("act_copy_path", "Copy path", "Path"),
                            ("act_copy_file", "Copy file", "Copy"),
                            ("act_terminal", "Open terminal here", "Term"),
                            ("act_reveal", "Show in Explorer", "Show"),
                            ("act_move", "Move to…", "Move"),
                            ("act_copy_to", "Copy to…", "Copy to"),
                            ("act_float", "Float (keep handy)", "Float"),
                            ("act_share", "Open with…", "Open")):
    REGISTRY.append(Op(id=_id, label=_label, glyph=_glyph, mode="action",
                       cats=frozenset({"*"}), ui=True, batch=True))


def _saving(src: Path, out: Path) -> str:
    a, b = src.stat().st_size, out.stat().st_size
    pct = (1 - b / a) * 100 if a else 0
    return f"{human_size(a)} → {human_size(b)} ({pct:.0f}% smaller)" if pct > 0 else \
        f"{human_size(a)} → {human_size(b)}"
